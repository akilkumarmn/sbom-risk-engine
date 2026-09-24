"""Formatting primitives for the Capstone-2 report (python-docx).

House style (matches the Capstone-1 report, with the requested sizes):
Times New Roman everywhere; every heading 14 pt bold; all body text,
tables and captions 12 pt; body justified at 1.5 line spacing; US Letter
with 1-inch margins; "Page X of Y" footer; REVA logo on front-matter pages.
"""
from __future__ import annotations

import textwrap
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402,F401
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402

FONT = "Times New Roman"
BODY = Pt(12)
HEAD = Pt(14)
BLACK = RGBColor(0, 0, 0)
CENTER, LEFT, JUSTIFY, RIGHT = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                                WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT)


# ---------------------------------------------------------------------------
# styles
# ---------------------------------------------------------------------------
def _force_font(el_rpr):
    rfonts = el_rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        el_rpr.insert(0, rfonts)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), FONT)


def _style(doc, name, size, bold=False, align=None, before=0, after=6, line=1.5, keep_next=False,
           base="Normal", italic=False):
    styles = doc.styles
    try:
        st = styles[name]
    except KeyError:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base]
    f = st.font
    f.name, f.size, f.bold, f.italic = FONT, size, bold, italic
    f.color.rgb = BLACK
    _force_font(st.element.get_or_add_rPr())
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    return st


def new_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    sec.footer_distance = Inches(0.5)
    # document defaults -> Times New Roman
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _force_font(rpr_default)
    _style(doc, "Normal", BODY, align=JUSTIFY, after=6, line=1.5)
    for lvl in (1, 2, 3):
        st = _style(doc, f"Heading {lvl}", HEAD, bold=True, align=CENTER if lvl == 1 else LEFT,
                    before=12 if lvl == 1 else 10, after=10 if lvl == 1 else 6, line=1.5, keep_next=True)
        st.font.italic = False
    _style(doc, "Figure Caption", BODY, bold=True, align=CENTER, before=4, after=12, line=1.15)
    _style(doc, "Table Caption", BODY, bold=True, align=CENTER, before=10, after=4, line=1.15, keep_next=True)
    _style(doc, "Table Text", BODY, align=LEFT, before=0, after=0, line=1.0)
    _style(doc, "Plain Center", BODY, align=CENTER, before=0, after=6, line=1.5)
    _style(doc, "Equation", BODY, align=CENTER, before=6, after=6, line=1.5, italic=True)
    _style(doc, "Code Text", BODY, align=LEFT, before=0, after=0, line=1.0)
    lb = _style(doc, "List Bullet", BODY, align=JUSTIFY, after=4, line=1.5)
    lb.paragraph_format.left_indent = Inches(0.35)
    for n in ("TOC 1", "TOC 2"):
        try:
            doc.styles[n]
        except KeyError:
            continue  # Word creates the built-in TOC styles itself (they inherit Times New Roman)
        _style(doc, n, BODY, after=2, line=1.15)
    # ask Word to refresh TOC / lists / page fields when the file is opened
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    after = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
             "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
             "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
             "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
    anchor = next((el for el in settings if el.tag.split("}")[1] in after), None)
    if anchor is not None:
        anchor.addprevious(upd)
    else:
        settings.append(upd)
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    _footer(doc)
    return doc


def _field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run._r.append(b)
    run2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run2._r.append(it)
    run3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run3._r.append(s)
    run4 = paragraph.add_run(placeholder)
    run4.font.name, run4.font.size = FONT, BODY
    run5 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run5._r.append(e)


def _footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True  # no page number on the cover
    p = sec.footer.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.alignment = RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Page ")
    r.font.name, r.font.size = FONT, BODY
    _field(p, "PAGE", "1")
    r = p.add_run(" of ")
    r.font.name, r.font.size = FONT, BODY
    _field(p, "NUMPAGES", "1")


# ---------------------------------------------------------------------------
# content helpers
# ---------------------------------------------------------------------------
class Report:
    def __init__(self, doc, logo: Path):
        self.doc = doc
        self.logo = logo
        self.fig_no: dict[int, int] = {}
        self.tab_no: dict[str, int] = {}
        self.figures: list[str] = []
        self.tables: list[str] = []

    # --- structure ---
    def page_break(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def logo_line(self):
        p = self.doc.add_paragraph(style="Plain Center")
        p.add_run().add_picture(str(self.logo), width=Inches(2.5))

    def h1(self, text, new_page=True, logo=False):
        if new_page:
            self.page_break()
        if logo:
            self.logo_line()
        return self.doc.add_paragraph(text, style="Heading 1")

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def p(self, *parts, style="Normal", align=None, italic=False):
        """parts: str, or (str, 'b'|'i'|'bi')."""
        para = self.doc.add_paragraph(style=style)
        if align is not None:
            para.alignment = align
        for part in parts:
            txt, fmt = (part, "") if isinstance(part, str) else part
            r = para.add_run(txt)
            r.bold = "b" in fmt
            r.italic = italic or "i" in fmt
        return para

    def center(self, text, bold=False, size=None, after=6):
        para = self.doc.add_paragraph(style="Plain Center")
        para.paragraph_format.space_after = Pt(after)
        r = para.add_run(text)
        r.bold = bold
        if size:
            r.font.size = size
        return para

    def bullets(self, items):
        for it in items:
            para = self.doc.add_paragraph(style="List Bullet")
            parts = it if isinstance(it, (list, tuple)) and not isinstance(it, str) else [it]
            for part in parts:
                txt, fmt = (part, "") if isinstance(part, str) else part
                r = para.add_run(txt)
                r.bold = "b" in fmt

    def numbered(self, items, fmt="{}."):
        """Manual numbering (restarts reliably in every list)."""
        for k, it in enumerate(items, 1):
            para = self.doc.add_paragraph(style="Normal")
            pf = para.paragraph_format
            pf.left_indent = Inches(0.4)
            pf.first_line_indent = Inches(-0.3)
            pf.space_after = Pt(4)
            label = fmt.format(k if "{}" in fmt else k)
            if fmt == "a)":
                label = f"{chr(96 + k)})"
            parts = it if isinstance(it, (list, tuple)) and not isinstance(it, str) else [it]
            para.add_run(label + "\t")
            pf.tab_stops.add_tab_stop(Inches(0.4))
            for part in parts:
                txt, f = (part, "") if isinstance(part, str) else part
                r = para.add_run(txt)
                r.bold = "b" in f

    def equation(self, text, number):
        para = self.doc.add_paragraph(style="Equation")
        para.add_run(text)
        r = para.add_run(f"    ({number})")
        r.italic = False
        return para

    def code(self, lines):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.columns[0].width = Inches(6.5)
        cell = t.cell(0, 0)
        cell.width = Inches(6.5)
        _shade(cell, "F2F2F2")
        first = True
        for line in lines:
            para = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            para.style = self.doc.styles["Code Text"]
            para.add_run(line)
        self.doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(2)

    # --- numbered figures / tables ---
    def figure(self, path: Path, chapter, caption, width=6.0):
        n = self.fig_no.get(chapter, 0) + 1
        self.fig_no[chapter] = n
        label = f"Fig. {chapter}.{n}"
        para = self.doc.add_paragraph(style="Plain Center")
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Inches(width))
        self.doc.add_paragraph(f"{label}: {caption}", style="Figure Caption")
        self.figures.append(label)
        return label

    def table(self, chapter, caption, header, rows, widths, bold_first_col=False, highlight_row=None):
        label = None
        if caption is not None:
            n = self.tab_no.get(str(chapter), 0) + 1
            self.tab_no[str(chapter)] = n
            label = f"Table {chapter}.{n}"
            self.doc.add_paragraph(f"{label}: {caption}", style="Table Caption")
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for j, w in enumerate(widths):
            t.columns[j].width = Inches(w)
        for i, row in enumerate([header] + rows):
            tr = t.rows[i]
            if i == 0:
                trpr = tr._tr.get_or_add_trPr()
                h = OxmlElement("w:tblHeader")
                h.set(qn("w:val"), "true")
                trpr.append(h)
            cant = OxmlElement("w:cantSplit")
            tr._tr.get_or_add_trPr().append(cant)
            for j, val in enumerate(row):
                cell = tr.cells[j]
                cell.width = Inches(widths[j])
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Table Text"]
                r = para.add_run(str(val))
                r.bold = i == 0 or (bold_first_col and j == 0) or (highlight_row is not None and i == highlight_row)
                if i == 0:
                    _shade(cell, "D9D9D9")
                    para.alignment = CENTER
                elif highlight_row is not None and i == highlight_row:
                    _shade(cell, "EDEDED")
        if len(rows) <= 12:  # keep short tables on one page (Word honours keep-with-next inside tables)
            for tr in t.rows[:-1]:
                for cell in tr.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next = True
        spacer = self.doc.add_paragraph(style="Normal")
        spacer.paragraph_format.space_after = Pt(4)
        if label:
            self.tables.append(label)
        return label

    def toc(self, instr, placeholder):
        para = self.doc.add_paragraph(style="Normal")
        _field(para, instr, placeholder)

    def sign_block(self, left, right, signature: Path | None = None):
        """Two-column borderless block, e.g. Place / Name and Date / Signature."""
        t = self.doc.add_table(rows=len(left), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j in range(2):
            t.columns[j].width = Inches(3.25)
        for i, (a, b) in enumerate(zip(left, right)):
            for j, txt in enumerate((a, b)):
                cell = t.cell(i, j)
                cell.width = Inches(3.25)
                para = cell.paragraphs[0]
                para.style = self.doc.styles["Table Text"]
                para.paragraph_format.space_after = Pt(6)
                para.add_run(txt)
                if signature is not None and Path(signature).exists() and j == 1 and i == len(left) - 1:
                    para.add_run("  ")
                    para.add_run().add_picture(str(signature), width=Inches(0.98))
        _no_borders(t)
        self.doc.add_paragraph(style="Normal")


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _no_borders(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    anchor = next((el for el in tblpr if el.tag.split("}")[1] in ("shd", "tblLayout", "tblCellMar", "tblLook",
                                                                   "tblCaption", "tblDescription")), None)
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tblpr.append(borders)


# ---------------------------------------------------------------------------
# diagrams (drawn, so they always match the implementation text)
# ---------------------------------------------------------------------------
INK, GREY, TEAL, ORANGE, LIGHT = "#222222", "#6B6B6B", "#1B7F74", "#C8651B", "#F2F4F6"


def _box(ax, x, y, w, h, title, body="", fc=LIGHT, ec="#9AA5B1", tc=INK, fs=9.5, wrap=28):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.06", fc=fc, ec=ec, lw=1.1))
    ax.text(x + w / 2, y + h - 0.13, title, ha="center", va="top", fontsize=fs, fontweight="bold", color=tc,
            family="serif")
    if body:
        ax.text(x + w / 2, y + h - 0.45, "\n".join(textwrap.wrap(body, wrap)), ha="center", va="top",
                fontsize=fs - 1.5, color=GREY if tc == INK else tc, family="serif", linespacing=1.25)


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=12, lw=1.3, color=GREY))


def draw_methodology(path: Path):
    steps = [("1. Data acquisition", "NVD, CISA KEV, FIRST EPSS, PoC-in-GitHub, Exploit-DB"),
             ("2. Feature engineering", "CVSS vector, CWE, vendor KEV rate, references, keywords, PoC within 30 days"),
             ("3. Model training", "Gradient boosting, time split 2019-2023 / 2024 / 2025, calibration"),
             ("4. Evaluation (Table 1)", "AUC-ROC, PR-AUC, precision@k, effort to cover 90%"),
             ("5. SBOM correlation", "CycloneDX SBOM + Trivy / Nessus / CSV scan, fan-in, asset context"),
             ("6. Context-aware scoring", "CVSS-only, proposal-stage formula, Formula + ML"),
             ("7. Time-frozen backtest", "Freeze 2025-01-01, answer key = KEV additions after it; ablation"),
             ("8. Deployment", "Dashboard on GitHub Pages, nightly retrain with gate, REST API")]
    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 4.6)
    ax.axis("off")
    w, h = 2.2, 1.65
    for k, (t, b) in enumerate(steps):
        row, col = divmod(k, 4)
        x = 0.2 + col * 2.45
        y = 2.65 - row * 2.35
        _box(ax, x, y, w, h, t, b, fc="#EAF4F2" if k < 4 else "#FBEFE5",
             ec=TEAL if k < 4 else ORANGE, wrap=26)
        if col < 3:
            _arrow(ax, x + w, y + h / 2, x + 2.45, y + h / 2)
    _arrow(ax, 0.2 + 3 * 2.45 + w / 2, 2.65, 0.2 + w / 2, 0.3 + h)
    ax.text(0.2, 4.52, "Model pipeline (offline, nightly)", fontsize=9, color=TEAL, family="serif", fontweight="bold",
            va="top")
    ax.text(0.2, 0.18, "Scoring, validation and deployment", fontsize=9, color=ORANGE, family="serif",
            fontweight="bold", va="top")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def draw_architecture(path: Path):
    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5.2)
    ax.axis("off")
    lanes = [(0.1, 2.6, "Public data sources"), (2.95, 3.3, "Nightly pipeline (GitHub Actions)"),
             (6.5, 3.4, "Delivery (GitHub Pages / REST API)")]
    for x, w, t in lanes:
        ax.add_patch(FancyBboxPatch((x, 0.15), w, 4.9, boxstyle="round,pad=0.02,rounding_size=0.08", fc="#FAFAFA",
                                    ec="#C9CED4", lw=1))
        ax.text(x + w / 2, 4.95, t, ha="center", va="top", fontsize=10, fontweight="bold", family="serif",
                color=INK)
    for k, (t, b) in enumerate([("NVD CVE feed", "fkie-cad mirror, CVSS, CWE, CPE, references"),
                                ("CISA KEV", "label; dateAdded"), ("FIRST EPSS", "current + freeze-date file"),
                                ("PoC sources", "PoC-in-GitHub, Exploit-DB")]):
        _box(ax, 0.3, 3.65 - k * 1.1, 2.2, 0.9, t, b, wrap=30)
    pipe = [("Feature builder", "label-free features, leakage guards"),
            ("KEV classifier", "gradient boosting, Platt calibration"),
            ("Export", "model_meta.json, per-CVE shards, KEV snapshot"),
            ("Backtest + gate", "Tables 2-3, tests, AUC > 0.80")]
    for k, (t, b) in enumerate(pipe):
        _box(ax, 3.15, 3.65 - k * 1.1, 2.9, 0.9, t, b, fc="#EAF4F2", ec=TEAL, wrap=36)
        if k:
            _arrow(ax, 4.6, 3.65 - (k - 1) * 1.1, 4.6, 3.65 - k * 1.1 + 0.9)
    dl = [("Uploads (stay in browser)", "CycloneDX SBOM, Trivy / Nessus / CSV scan, asset context"),
          ("Correlation + enrichment", "bom-ref / purl / name@version, fan-in, services"),
          ("Live intel", "EPSS API and KEV JSON, 24 h cache, snapshot fallback"),
          ("Ranked output", "CVSS-only | Formula | Formula + ML, explanations")]
    for k, (t, b) in enumerate(dl):
        _box(ax, 6.7, 3.65 - k * 1.1, 3.0, 0.9, t, b, fc="#FBEFE5", ec=ORANGE, wrap=38)
        if k:
            _arrow(ax, 8.2, 3.65 - (k - 1) * 1.1, 8.2, 3.65 - k * 1.1 + 0.9)
    _arrow(ax, 2.5, 2.55, 3.15, 2.55)
    _arrow(ax, 6.05, 1.45, 6.7, 1.45)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


def draw_mlops(path: Path):
    steps = ["Fetch feeds", "Build features", "Train (Table 1)", "Export shards", "Backtest (Tables 2-3)",
             "Tests + AUC gate", "Commit + deploy"]
    fig, ax = plt.subplots(figsize=(10, 2.3))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2.3)
    ax.axis("off")
    w = 1.25
    for k, s in enumerate(steps):
        x = 0.1 + k * 1.42
        gate = s.startswith("Tests")
        _box(ax, x, 0.75, w, 0.95, "\n".join(textwrap.wrap(s, 12)), fc="#FDECEC" if gate else "#EAF4F2",
             ec="#C0392B" if gate else TEAL, fs=9)
        if k:
            _arrow(ax, x - 0.17, 1.22, x, 1.22)
    ax.text(5.0, 2.2, "retrain.yml: cron '30 21 * * *' (03:00 IST) and manual dispatch", ha="center", va="top",
            fontsize=9.5, family="serif", color=INK)
    ax.text(5.0, 0.45, "A failed test or gate stops the run: nothing is committed and the live site keeps the last "
            "good model.", ha="center", va="top", fontsize=9, family="serif", color=GREY)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)

"""Build the Cap-2 results deck from the REVA template and the pipeline output.

    python -m scripts.build_deck --date "September 2026" --url https://<user>.github.io/sbom-risk-engine/

Every number on the results slides is read from site/model_meta.json and
site/backtest.json, so re-running this after the real pipeline run replaces the
synthetic numbers. Result slides carry a red "SYNTHETIC TEST DATA" stamp while
those files say data_source == synthetic-test-fixture.

docs/deck/template.pptx is the Cap-1 deck with the structural fixes already
applied: the duplicate methodology slide removed and seven content slides added.
"""
from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from engine import config

INK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
PANEL = RGBColor(0xF2, 0xF2, 0xF2)
LINE = RGBColor(0xD0, 0xD0, 0xD0)
ORANGE = RGBColor(0xE0, 0x7B, 0x24)
TEAL = RGBColor(0x1B, 0x99, 0x8B)
RED = RGBColor(0xD6, 0x3B, 0x40)
PURPLE = RGBColor(0x6A, 0x5A, 0xCD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------------------------------------------------------------------------
# primitives
# ---------------------------------------------------------------------------
def text(slide, x, y, w, h, paras, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         margin=0.05):
    """paras: list of str or (str, dict(size,bold,color,italic)) or list-of-runs [(str, dict), ...]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            s, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = s
            f = run.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
        if isinstance(p, tuple) and p[1].get("space_after") is not None:
            para.space_after = Pt(p[1]["space_after"])
        else:
            para.space_after = Pt(4)
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, title, body, fill=PANEL, title_color=INK, size=12, badge=None, badge_fill=ORANGE):
    box(slide, x, y, w, h, fill)
    tx = x + 0.2
    if badge:
        c = box(slide, x + 0.2, y + 0.2, 0.42, 0.42, badge_fill, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = badge
        r.font.size, r.font.bold, r.font.name = Pt(12), True, FONT
        r.font.color.rgb = WHITE
        tx = x + 0.75
    text(slide, tx, y + 0.15, w - (tx - x) - 0.15, 0.7, [(title, {"bold": True, "size": size + 3, "color": title_color})],
         anchor=MSO_ANCHOR.MIDDLE if badge else MSO_ANCHOR.TOP)
    text(slide, x + 0.2, y + 0.95, w - 0.4, h - 1.05, body, size=size, color=INK)


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, f"{A}tailEnd")
    tail.set("type", "triangle")
    return c


def table(slide, x, y, w, rows, col_w, size=11, header_fill=RGBColor(0x40, 0x40, 0x40), highlight=None, row_h=0.34):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows)).table
    for j, cw in enumerate(col_w):
        gt.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        gt.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = gt.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif highlight is not None and i == highlight:
                cell.fill.fore_color.rgb = RGBColor(0xE3, 0xF4, 0xF1)
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else PANEL
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or (highlight is not None and i == highlight)
            r.font.color.rgb = WHITE if i == 0 else INK
    return gt


def stamp(slide, synthetic):
    if not synthetic:
        return
    s = box(slide, 0.35, 1.2, 4.7, 0.34, RGBColor(0xFD, 0xEC, 0xEC), line=RED)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = "SYNTHETIC TEST DATA — re-run after the real pipeline"
    r.font.size, r.font.bold, r.font.name = Pt(10.5), True, FONT
    r.font.color.rgb = RED


def pct(x, d=1):
    return "n/a" if x is None else f"{100 * x:.{d}f}%"


def f3(x):
    return "n/a" if x is None else f"{x:.3f}"


# ---------------------------------------------------------------------------
# template housekeeping
# ---------------------------------------------------------------------------
KEEP = ("Title", "Date Placeholder", "Footer Placeholder", "Slide Number Placeholder")


def clear_body(slide):
    for sh in list(slide.shapes):
        if sh.is_placeholder and sh.placeholder_format.type in (1, 3, 13, 15, 16):
            continue
        sh._element.getparent().remove(sh._element)


def set_placeholder_text(shape, value):
    """Replace a placeholder's content (e.g. a date field) by plain text, keeping the first run's formatting."""
    txBody = shape._element.txBody
    p = txBody.find(f"{A}p")
    rpr = None
    for el in p:
        if el.tag in (f"{A}r", f"{A}fld"):
            r0 = el.find(f"{A}rPr")
            rpr = copy.deepcopy(r0) if r0 is not None else None
            break
    for el in list(p):
        if el.tag in (f"{A}r", f"{A}fld", f"{A}br"):
            p.remove(el)
    for extra in txBody.findall(f"{A}p")[1:]:
        txBody.remove(extra)
    r = etree.SubElement(p, f"{A}r")
    if rpr is not None:
        r.append(rpr)
    t = etree.SubElement(r, f"{A}t")
    t.text = value
    end = p.find(f"{A}endParaRPr")
    if end is not None:
        p.remove(end)
        p.append(end)


def footer_fix(prs, date):
    n = len(prs.slides) - 1  # the closing slide carries no number
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.is_placeholder:
                continue
            t = sh.placeholder_format.type
            if t == 16:
                set_placeholder_text(sh, date)
            elif t == 13 and i > 1:
                set_placeholder_text(sh, f"{i}/{n}")


def set_title(slide, value):
    title = slide.shapes.title
    run = title.text_frame.paragraphs[0].runs[0]
    run.text = value
    for r in title.text_frame.paragraphs[0].runs[1:]:
        r.text = ""


# ---------------------------------------------------------------------------
# slides
# ---------------------------------------------------------------------------
def s_title(slide, date):
    sub = next(sh for sh in slide.shapes if sh.is_placeholder and sh.placeholder_format.type == 4)
    for p in sub.text_frame.paragraphs:
        if p.runs and p.runs[0].text.startswith("Date"):
            p.runs[0].text = f"Date: {date}"
            for r in p.runs[1:]:
                r.text = ""


def s_agenda(slide):
    clear_body(slide)
    items = [("Introduction", "Background | Why this study"), ("Literature Review", "Seminal works | Research gap"),
             ("Problem Statement", "Why CVSS alone fails"), ("Project Objectives", "What Cap-2 had to prove"),
             ("Architecture & Methodology", "Pipeline | Risk formula"),
             ("Data & Model", "Sources | Features | Table 1"), ("Backtest & Ablation", "Tables 2 and 3"),
             ("Live Demo & MLOps", "Dashboard | Nightly retrain | API"),
             ("Limitations & Future Work", "What the numbers do not show"), ("References", "Papers | Data sources")]
    for k, (t, sub) in enumerate(items):
        col, row = divmod(k, 5)
        x, y = 0.9 + col * 6.1, 1.45 + row * 1.05
        c = box(slide, x, y + 0.05, 0.62, 0.62, ORANGE if col == 0 else RGBColor(0x40, 0x40, 0x40), shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"{k + 1:02d}"
        r.font.size, r.font.bold, r.font.name = Pt(15), True, FONT
        r.font.color.rgb = WHITE
        text(slide, x + 0.8, y, 4.9, 0.75, [(t, {"bold": True, "size": 17}), (sub, {"size": 12, "color": MUTED})])


def s_fix_typos(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("which fails", "which fail").replace("in efficiencies", "inefficiencies").replace("exposing system", "exposing systems")


def s_objectives(slide):
    clear_body(slide)
    set_title(slide, "Project Objectives")
    text(slide, 0.6, 1.3, 12.1, 0.5, [("Cap-1 proposed a context-aware score; Cap-2 had to make the \"AI\" real, "
                                        "measure it against CVSS and deploy it.", {"size": 15, "color": MUTED})])
    cards = [
        ("1", "Learned exploitation likelihood",
         ["Gradient-boosted classifier trained on NVD CVEs with CISA KEV as the label.",
          "Time split: train 2019–2023, validate 2024, test 2025.",
          "Compared with CVSS-only and EPSS-only baselines (Table 1)."]),
        ("2", "Context-aware ranking",
         ["Five signals: model probability, EPSS, KEV, SBOM fan-in with asset criticality, CVSS scope.",
          "Three rankings side by side: CVSS-only, Cap-1 formula, Formula + ML.",
          "Business context from an optional asset-context file."]),
        ("3", "Evidence and deployment",
         ["Time-frozen backtest: patch effort to cover 90% of later-exploited CVEs (Tables 2–3).",
          "Nightly retrain with a regression gate, public dashboard, REST API."]),
    ]
    for i, (n, t, body) in enumerate(cards):
        card(slide, 0.6 + i * 4.1, 2.05, 3.85, 4.3, t, [(b, {"space_after": 10, "size": 14}) for b in body], badge=n, size=13)


def s_architecture(slide):
    clear_body(slide)
    set_title(slide, "Architecture and Methodology")
    # three swimlanes
    lanes = [(0.5, "Public sources"), (4.05, "Nightly pipeline (GitHub Actions)"), (8.55, "Dashboard (GitHub Pages) / API")]
    widths = [3.2, 4.1, 4.3]
    for (x, label), w in zip(lanes, widths):
        box(slide, x, 1.3, w, 3.75, PANEL)
        text(slide, x + 0.15, 1.38, w - 0.3, 0.35, [(label, {"bold": True, "size": 12, "color": MUTED})])
    for k, s in enumerate(["NVD CVE feed (fkie-cad mirror)", "CISA KEV catalogue", "FIRST EPSS (current + 2025-01-01)",
                           "PoC-in-GitHub + Exploit-DB"]):
        b = box(slide, 0.7, 1.85 + k * 0.78, 2.8, 0.6, WHITE, line=LINE)
        text(slide, 0.75, 1.9 + k * 0.78, 2.7, 0.5, [(s, {"size": 11})], anchor=MSO_ANCHOR.MIDDLE)
    steps = [("Feature builder", "CVSS vector, CWE, vendor KEV rate, references, keywords, PoC ≤30 days"),
             ("KEV classifier", "LightGBM, time split, Platt-calibrated"),
             ("Export + backtest", "per-CVE probability + top-3 reasons; Tables 1–3; AUC gate > 0.80")]
    for k, (t, d) in enumerate(steps):
        box(slide, 4.25, 1.85 + k * 1.03, 3.7, 0.85, WHITE, line=LINE)
        text(slide, 4.33, 1.88 + k * 1.03, 3.55, 0.8, [(t, {"bold": True, "size": 12}), (d, {"size": 10, "color": MUTED})])
    for k, (t, d) in enumerate([("Inputs (stay in the browser)", "CycloneDX SBOM · Trivy / Nessus / CSV · asset-context.json"),
                                ("SBOM enrichment", "correlation by bom-ref / purl; dependency depth and fan-in"),
                                ("Live intel", "EPSS and KEV fetched live, cached 24 h, dated snapshot fallback")]):
        box(slide, 8.75, 1.85 + k * 1.03, 3.9, 0.85, WHITE, line=LINE)
        text(slide, 8.83, 1.88 + k * 1.03, 3.75, 0.8, [(t, {"bold": True, "size": 12}), (d, {"size": 10, "color": MUTED})])
    arrow(slide, 3.72, 3.2, 4.02, 3.2)
    arrow(slide, 7.93, 3.2, 8.53, 3.2)
    # formula
    box(slide, 0.5, 5.2, 12.35, 1.55, RGBColor(0x40, 0x40, 0x40))
    text(slide, 0.7, 5.28, 12.0, 1.45, [
        ("Formula + ML risk score", {"bold": True, "size": 13, "color": ORANGE}),
        ("Risk = 100 × (0.5·P_model + 0.3·EPSS + 0.2·KEV) × (0.5·CVSS impact/6 + 0.5·fan-in/max fan-in) × Crit_asset × Scope",
         {"size": 14, "color": WHITE, "bold": True}),
        ("Crit_asset = 1.5 / 1.25 / 1.0 for tier 1 / 2 / 3, ×1.2 if internet-facing · Scope = 1.15 if CVSS scope changed · "
         "SBOM depth and fan-in are not reachability", {"size": 11, "color": RGBColor(0xDD, 0xDD, 0xDD)}),
    ])


def s_data(slide, meta):
    clear_body(slide)
    set_title(slide, "Dataset, Features and Split")
    c = meta["counts"]
    total = sum(v["n"] for v in c.values())
    kev = sum(v["kev"] for v in c.values())
    rows = [["Source", "Used for"],
            ["NVD CVE feed (fkie-cad mirror)", "CVSS vector, CWE, CPE vendor, references, dates"],
            ["CISA KEV catalogue", "label; dateAdded rebuilds KEV at any past date"],
            ["FIRST EPSS", "baseline, model variant, backtest input"],
            ["PoC-in-GitHub, Exploit-DB", "first public-PoC dates"]]
    table(slide, 0.5, 1.7, 6.6, rows, [2.6, 4.0], size=11, row_h=0.42)
    text(slide, 0.5, 4.0, 6.6, 2.7, [
        ("Features (knowable at publication)", {"bold": True, "size": 14}),
        ("CVSS base + one-hot vector · top-30 CWE · vendor/product KEV rate (target-encoded on training years, "
         "out-of-fold) · reference counts and exploit-site domains · description keywords (RCE, unauthenticated, "
         "deserialization, SSRF, auth bypass) · public PoC within 30 days of publication", {"size": 12}),
        ("Leakage guards: KEV-derived NVD fields and KEV-catalogue links are ignored; tested with planted traps.",
         {"size": 12, "color": MUTED}),
    ])
    # split timeline
    x0 = 7.6
    text(slide, x0, 1.7, 5.2, 0.4, [("Time-based split (NVD published date)", {"bold": True, "size": 14})])
    segs = [("Train", "2019–2023", c["train"], 2.2, RGBColor(0x40, 0x40, 0x40)),
            ("Validate", str(meta["split"]["val_year"]), c["val"], 1.35, MUTED),
            ("Test", str(meta["split"]["test_year"]), c["test"], 1.35, ORANGE)]
    x = x0
    for name, yrs, cnt, w, col in segs:
        b = box(slide, x, 2.25, w, 0.75, col, shape=MSO_SHAPE.RECTANGLE)
        text(slide, x, 2.27, w, 0.72, [(f"{name} {yrs}", {"bold": True, "size": 11, "color": WHITE}),
                                       (f"{cnt['n']:,} CVEs · {cnt['kev']} KEV", {"size": 9.5, "color": WHITE})],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w + 0.08
    for k, (big, lab) in enumerate([(f"{total:,}", "CVEs in the split"), (f"{kev:,}", "in CISA KEV"),
                                    (pct(kev / total if total else None), "positive rate")]):
        text(slide, x0 + k * 1.75, 3.4, 1.7, 1.2, [(big, {"bold": True, "size": 26, "color": ORANGE}),
                                                     (lab, {"size": 11, "color": MUTED})])
    text(slide, x0, 4.75, 5.2, 1.9, [
        ("Why not a random split?", {"bold": True, "size": 13}),
        ("A random split lets 2025 CVEs learn from their own vendors' later history. Splitting by date mimics "
         "deployment: the model only ever scores CVEs newer than everything it trained on.", {"size": 12}),
    ])


def s_model(slide, meta, synthetic):
    clear_body(slide)
    set_title(slide, "Model Results — Table 1")
    stamp(slide, synthetic)
    t = {r["ranker"]: r for r in meta["table1"]}
    order = [("cvss", "CVSS base score"), ("epss", "EPSS (FIRST)*"), ("model_no_epss", "Our model, without EPSS"),
             ("model_with_epss", "Our model, with EPSS*"), ("logreg_no_epss", "Logistic regression")]
    rows = [["Ranker", "AUC-ROC", "PR-AUC", "P@100", "R@1000", "Effort 90%"]]
    for k, name in order:
        r = t[k]
        rows.append([name, f3(r["auc_roc"]), f3(r["pr_auc"]), f3(r["precision_at_100"]), f3(r["recall_at_1000"]),
                     pct(r["effort_to_cover_90"])])
    c = meta["counts"]["test"]
    text(slide, 0.5, 1.65, 7.2, 0.4, [(f"{meta['split']['test_year']} test year: n = {c['n']:,} CVEs, "
                                        f"k = {c['kev']} in KEV ({pct(c['kev_rate'])})", {"size": 12, "color": MUTED})])
    table(slide, 0.5, 2.05, 7.2, rows, [2.55, 0.9, 0.9, 0.9, 0.95, 1.0], size=11, highlight=3, row_h=0.4)
    text(slide, 0.5, 4.6, 7.2, 2.2, [
        (f"Learner: {meta['algo']} · Platt-calibrated on {meta['split']['val_year']} · metrics are expected values "
         "under random tie-breaking (CVSS ties thousands of CVEs).", {"size": 11, "color": MUTED}),
        ("* EPSS rows use today's EPSS file, which already reflects post-publication exploitation: an optimistic "
         "upper bound. The time-frozen comparison is the backtest.", {"size": 11, "color": MUTED}),
        ("Highlighted row = the project's own contribution, independent of EPSS.", {"size": 11, "color": MUTED}),
    ])
    rows_imp = (meta.get("shap_mean_abs") or meta["importance"])[:10][::-1]
    cd = CategoryChartData()
    cd.categories = [r["label"] for r in rows_imp]
    cd.add_series("importance", [round(r["value"], 4) for r in rows_imp])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(7.95), Inches(1.6), Inches(4.95), Inches(5.2), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    kind = "mean |SHAP|" if meta.get("shap_mean_abs") else meta["importance_kind"]
    ch.chart_title.text_frame.text = f"Top-10 features ({kind})"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.plots[0].series[0].format.fill.solid()
    ch.plots[0].series[0].format.fill.fore_color.rgb = TEAL
    ch.plots[0].gap_width = 60
    ch.category_axis.tick_labels.font.size = Pt(9)
    ch.value_axis.tick_labels.font.size = Pt(9)
    ch.value_axis.has_major_gridlines = False


def _coverage_chart(slide, case, x, y, w, h, title):
    cd = XyChartData()
    names = {"cvss": ("CVSS-only", RED), "legacy": ("Formula (Cap-1)", ORANGE), "ml": ("Formula + ML", TEAL),
             "epss": ("EPSS-only", PURPLE)}
    for k, (label, _) in names.items():
        cv = case["coverage"].get(k)
        if not cv:
            continue
        s = cd.add_series(label)
        for px, py in zip(cv["x"], cv["y"]):
            s.add_data_point(px, py)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(9)
    for s, (k, (label, col)) in zip(ch.plots[0].series, [(k, v) for k, v in names.items() if case["coverage"].get(k)]):
        s.format.line.color.rgb = col
        s.format.line.width = Pt(2.75 if k == "ml" else 1.5)
        s.smooth = False
    for ax in (ch.category_axis, ch.value_axis):
        ax.minimum_scale, ax.maximum_scale = 0, 1
        ax.tick_labels.font.size = Pt(9)
        ax.tick_labels.number_format = '0%'
        ax.tick_labels.number_format_is_linked = False
        ax.has_major_gridlines = ax is ch.value_axis
        if ax.has_major_gridlines:
            ax.major_gridlines.format.line.color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
    ch.category_axis.has_title = True
    ch.category_axis.axis_title.text_frame.text = "findings patched (ranked order)"
    ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    ch.value_axis.has_title = True
    ch.value_axis.axis_title.text_frame.text = "later-exploited covered"
    ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(9)


def s_backtest(slide, bt, synthetic):
    clear_body(slide)
    set_title(slide, "Backtest — Table 2")
    stamp(slide, synthetic)
    case = bt["cases"][bt["primary_case"]]
    h = bt["headline"]
    hp = bt["headline_population"]
    box(slide, 0.5, 1.7, 4.0, 2.05, RGBColor(0x40, 0x40, 0x40))
    text(slide, 0.65, 1.78, 3.7, 1.95, [
        ([(pct(h["ours"], 0), {"size": 34, "bold": True, "color": TEAL}),
          ("  vs  ", {"size": 16, "color": WHITE}),
          (pct(h["cvss"], 0), {"size": 34, "bold": True, "color": RGBColor(0xF0, 0x6A, 0x6E)})]),
        ("of findings patched to cover 90% of the CVEs added to KEV after the freeze — Formula + ML vs CVSS-only",
         {"size": 11, "color": WHITE}),
    ])
    text(slide, 0.5, 3.85, 4.0, 1.0, [
        (f"Freeze {bt['freeze_date']} · {case.get('app', case['name'])}: N = {case['n_findings']} findings, "
         f"M = {case['later_exploited']} later-exploited, {case['known_exploited_at_T']} already in KEV.",
         {"size": 11, "color": MUTED}),
        (f"Whole population before T (N = {hp['n_findings']:,}, M = {hp['later_exploited']}): "
         f"{pct(hp['ours'])} vs {pct(hp['cvss'])}.", {"size": 11, "color": MUTED}),
    ])
    rows = [["Ranking method", "Mean rank", "P@10", "Effort 90%"]]
    names = {"cvss": "CVSS-only", "legacy": "Formula (Cap-1)", "ml": "Formula + ML", "epss": "EPSS-only (ref.)"}
    t = {r["ranker"]: r for r in case["table2"]}
    for k, n in names.items():
        r = t[k]
        rows.append([n, "n/a" if r["mean_rank_later_exploited"] is None else f"{r['mean_rank_later_exploited']:.1f}",
                     f3(r["precision_at_10"]), pct(r["effort_to_cover_90"])])
    table(slide, 0.5, 4.95, 4.0, rows, [1.6, 0.85, 0.7, 0.85], size=10.5, highlight=3, row_h=0.34)
    _coverage_chart(slide, case, 4.8, 1.6, 8.1, 5.2, f"Coverage curve, frozen {bt['freeze_date']}: {case['name']}")


def s_ablation(slide, bt, synthetic):
    clear_body(slide)
    set_title(slide, "Ablation — Table 3")
    stamp(slide, synthetic)
    prim = bt["primary_case"]
    cases = [prim, "population"]
    abl = {c: {r["removed"]: r for r in bt["cases"][c]["ablation"]} for c in cases}
    rows = [["Signal removed", f"{prim}: effort 90%", "change", "population: effort 90%", "change"]]
    for key, r in abl[prim].items():
        row = [r["name"]]
        for c in cases:
            a = abl[c][key]
            row += [pct(a["effort_to_cover_90"]),
                    "—" if a["delta_vs_full"] is None else f"{100 * a['delta_vs_full']:+.1f} pp"]
        rows.append(row)
    table(slide, 0.5, 1.7, 7.9, rows, [2.5, 1.45, 0.95, 2.05, 0.95], size=11, highlight=1, row_h=0.42)
    lines = []
    for key, r in abl["population"].items():
        if key == "none":
            continue
        d = r["delta_vs_full"] if r["delta_vs_full"] is not None else abl[prim][key]["delta_vs_full"]
        where = "population" if r["delta_vs_full"] is not None else prim
        if d is None:
            continue
        verdict = "helps" if d > 0.005 else "hurts" if d < -0.005 else "no measurable effect"
        lines.append((f"{r['name']}: {verdict} ({100 * d:+.1f} pp on {where})", {"size": 12}))
    box(slide, 8.7, 1.7, 4.15, 5.0, PANEL)
    text(slide, 8.85, 1.8, 3.9, 4.8, [("Reading the table", {"bold": True, "size": 14})] + lines + [
        ("Positive change = removing the signal made the ranking worse. A signal that hurts is reported, "
         "not hidden.", {"size": 11, "color": MUTED})])
    text(slide, 0.5, 5.0, 7.9, 1.6, [
        ("Only the Likelihood weights (0.5 / 0.3 / 0.2) are hand-set; this table is the evidence for keeping or "
         "dropping each term. Fan-in and asset criticality only exist in the SBOM case.", {"size": 12, "color": MUTED})])


def s_demo(slide, fig, url):
    clear_body(slide)
    set_title(slide, "Live Demo — Dashboard")
    if fig.exists():
        from PIL import Image
        w_px, h_px = Image.open(fig).size
        h = 5.25
        w = min(8.2, h * w_px / h_px)
        slide.shapes.add_picture(str(fig), Inches(0.5), Inches(1.45), Inches(w), Inches(w * h_px / w_px))
    pts = [("Three-way toggle", "CVSS-only → Formula → Formula + ML: the ranking moves twice."),
           ("ML signal card", "Test AUC read from model_meta.json, never typed by hand."),
           ("Per-finding breakdown", "P_model and its top-3 feature contributions, formula terms, services."),
           ("Live threat intel", "EPSS and KEV fetched live with a dated snapshot fallback; only CVE IDs leave the browser."),
           ("Asset context upload", "Tier and exposure multiply Impact.")]
    y = 1.45
    for t, d in pts:
        text(slide, 8.95, y, 3.95, 0.95, [(t, {"bold": True, "size": 13}), (d, {"size": 11, "color": MUTED})])
        y += 0.97
    text(slide, 8.95, y + 0.05, 3.95, 0.5, [(url, {"size": 11, "color": ORANGE, "bold": True})])


def s_mlops(slide):
    clear_body(slide)
    set_title(slide, "MLOps — Nightly Retrain and API")
    steps = ["Fetch feeds", "Build features", "Train + Table 1", "Export shards", "Backtest", "Tests + gate", "Commit + deploy"]
    w, gap, x0, y = 1.62, 0.17, 0.5, 1.7
    for k, s in enumerate(steps):
        x = x0 + k * (w + gap)
        col = RED if s == "Tests + gate" else RGBColor(0x40, 0x40, 0x40)
        b = box(slide, x, y, w, 0.85, col)
        text(slide, x, y, w, 0.85, [(s, {"bold": True, "size": 12, "color": WHITE})], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        if k:
            arrow(slide, x - gap + 0.01, y + 0.42, x - 0.01, y + 0.42)
    text(slide, 0.5, 2.65, 12.3, 0.4, [("retrain.yml · cron 30 21 * * * (03:00 IST) · the Pages deploy is called "
                                        "directly after the commit", {"size": 11, "color": MUTED})])
    card(slide, 0.5, 3.25, 4.0, 2.9, "Regression gate",
         [("The nightly model is published only if test AUC > 0.80 on real data, the backtest exists and the "
           "Python/JavaScript parity tests pass.", {"size": 13, "space_after": 8}),
          ("A failed gate keeps yesterday's model online.", {"size": 12, "color": MUTED})], badge="✓", badge_fill=RED)
    card(slide, 4.7, 3.25, 4.0, 2.9, "One formula, two engines",
         [("engine/scoring.py (backtest, API) and site/engine.js (browser) are checked for identical scores on "
           "every push (tests.yml).", {"size": 12, "space_after": 8}),
          ("Leakage traps and the backtest answer key are tested too.", {"size": 12, "color": MUTED})],
         badge="=", badge_fill=TEAL)
    card(slide, 8.9, 3.25, 3.95, 2.9, "REST API (FastAPI + Docker)",
         [("POST /score: SBOM + scan (+ asset context) → ranked list with all three scores and top features.",
           {"size": 12, "space_after": 8}),
          ("docker run -p 8080:8080 sbom-risk-engine → Swagger UI at /docs", {"size": 11, "color": MUTED})],
         badge="{}", badge_fill=ORANGE)


def s_outcomes(slide, bt, meta, synthetic):
    clear_body(slide)
    set_title(slide, "Outcomes against the Cap-1 Expectations")
    stamp(slide, synthetic)
    h = bt["headline"]
    rows = [["Cap-1 expected result", "What Cap-2 measured", "Evidence"],
            ["Smarter identification of prioritized vulnerabilities",
             f"Model without EPSS: test AUC {f3(meta['test_auc'])} vs CVSS {f3({r['ranker']: r for r in meta['table1']}['cvss']['auc_roc'])}",
             "Table 1"],
            ["Reduction in patch workload",
             f"Effort to cover 90% of later-exploited CVEs: {pct(h['ours'])} (ours) vs {pct(h['cvss'])} (CVSS)",
             "Table 2"],
            ["Reduction in MTTR", "Not measurable without remediation timestamps — dropped as a claim",
             "Limitations"],
            ["Improved risk visibility", "Per-finding breakdown, top features, live intel badges, three-way toggle",
             "Dashboard"]]
    table(slide, 0.5, 1.75, 12.35, rows, [3.6, 6.9, 1.85], size=12, row_h=0.62)
    text(slide, 0.5, 5.15, 12.35, 1.4, [
        ("Rule used for this deck: every claim points at a number in Tables 1–3 or at the running system.",
         {"size": 13, "bold": True}),
        ("Numbers are read from site/model_meta.json and site/backtest.json by scripts/build_deck.py.",
         {"size": 11, "color": MUTED})])


def s_limits(slide):
    clear_body(slide)
    set_title(slide, "Limitations and Future Work")
    lim = ["KEV is an incomplete, US-federal-biased label; \"not in KEV\" is not \"never exploited\".",
           "Low prevalence limits precision; PR-AUC and precision@k are the honest metrics.",
           "EPSS already encodes much of the signal; the model without EPSS is the independent contribution.",
           "Backtest CVSS is today's NVD value; one SBOM yields few later-exploited CVEs (population case added).",
           "SBOM depth and fan-in are not reachability; asset context is user-supplied, not discovered.",
           "Nessus findings without a CPE are skipped (counted in the log)."]
    fut = ["Call-graph reachability analysis for Java dependencies.",
           "VEX ingestion to suppress not-affected findings.",
           "Learn the Likelihood weights end to end from backtest outcomes.",
           "Relabel with organisation-specific incident data."]
    box(slide, 0.5, 1.45, 7.4, 5.3, PANEL)
    text(slide, 0.7, 1.55, 7.0, 5.1, [("Limitations", {"bold": True, "size": 16})] +
         [(f"–  {s}", {"size": 14, "space_after": 10}) for s in lim])
    box(slide, 8.15, 1.45, 4.7, 5.3, RGBColor(0x40, 0x40, 0x40))
    text(slide, 8.35, 1.55, 4.3, 5.1, [("Future work", {"bold": True, "size": 16, "color": ORANGE})] +
         [(f"–  {s}", {"size": 14, "color": WHITE, "space_after": 10}) for s in fut])


def s_references(slide):
    clear_body(slide)
    set_title(slide, "References")
    refs = [
        ("A Statistical Approach to Severity Aware Vulnerability Prioritization, IEEE, 2024.",
         "https://ieeexplore-ieee-org-reva.knimbus.com/document/10815757"),
        ("User Vulnerabilities in AI-Driven Systems: Current Cybersecurity Threat Dynamics and Malicious Exploits in "
         "Supply Chain Management and Project Management, IEEE, 2024.",
         "https://ieeexplore-ieee-org-reva.knimbus.com/document/10459480"),
        ("ERS0: Enhancing Military Cybersecurity with AI-Driven SBOM for Firmware Vulnerability Detection and Asset "
         "Management, IEEE, 2024.", "https://ieeexplore-ieee-org-reva.knimbus.com/document/10685598"),
        ("Supply Chain Risk Analysis Via SBOM Data Enrichment, IEEE, 2023.",
         "https://ieeexplore-ieee-org-reva.knimbus.com/document/11014830"),
        ("J. Jacobs, S. Romanosky, B. Edwards, I. Adjerid, M. Roytman, \"Exploit Prediction Scoring System (EPSS)\", "
         "Digital Threats: Research and Practice, 2(3), 2021.", "https://www.first.org/epss/"),
        ("G. Ke et al., \"LightGBM: A Highly Efficient Gradient Boosting Decision Tree\", NeurIPS, 2017.", ""),
        ("CISA, Known Exploited Vulnerabilities Catalog.", "https://www.cisa.gov/known-exploited-vulnerabilities-catalog"),
        ("OWASP CycloneDX Bill of Materials Standard.", "https://cyclonedx.org/specification/overview/"),
    ]
    paras = []
    for k, (t, u) in enumerate(refs, 1):
        runs = [(f"{k}.  {t}", {"size": 13})]
        if u:
            runs.append((f"  {u}", {"size": 11, "color": TEAL}))
        paras.append(runs)
    tb = text(slide, 0.6, 1.4, 12.1, 5.3, paras)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(9)


# ---------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", type=Path, default=config.ROOT / "docs" / "deck" / "template.pptx")
    ap.add_argument("--out", type=Path, default=config.ROOT / "docs" / "deck" / "Cap2_SBOM_Risk_Engine_results.pptx")
    ap.add_argument("--date", default="September 2026", help="one date for the title slide and every footer")
    ap.add_argument("--url", default="https://<github-user>.github.io/sbom-risk-engine/")
    args = ap.parse_args(argv)
    meta = json.loads((config.SITE_DIR / "model_meta.json").read_text())
    bt = json.loads((config.SITE_DIR / "backtest.json").read_text())
    synthetic = "synthetic-test-fixture" in (meta.get("data_source"), bt.get("data_source"))

    prs = Presentation(str(args.template))
    sl = prs.slides
    assert len(sl) == 17, "template must have 17 slides"
    s_title(sl[0], args.date)
    s_agenda(sl[1])
    s_fix_typos(sl[2])
    s_objectives(sl[5])
    s_architecture(sl[6])
    s_data(sl[7], meta)
    s_model(sl[8], meta, synthetic)
    s_backtest(sl[9], bt, synthetic)
    s_ablation(sl[10], bt, synthetic)
    s_demo(sl[11], config.FIGURES_DIR / "dashboard.png", args.url)
    s_mlops(sl[12])
    s_outcomes(sl[13], bt, meta, synthetic)
    s_limits(sl[14])
    s_references(sl[15])
    footer_fix(prs, args.date)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    print(f"wrote {args.out} ({len(sl)} slides){' [SYNTHETIC numbers]' if synthetic else ''}")


if __name__ == "__main__":
    main()
